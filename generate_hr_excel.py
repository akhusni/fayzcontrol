import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

wb = openpyxl.Workbook()

# Sheet 1: Xodimlar Ro'yxati
ws = wb.active
ws.title = "Xodimlar Ro'yxati"
ws.views.sheetView[0].showGridLines = True

# Title banner
ws.merge_cells('A1:N1')
title_cell = ws['A1']
title_cell.value = "FAYZ MEDICAL HOUSE — XODIMLAR MA'LUMOTLARINI YIG'ISH JADVALI (HR & EMR)"
title_cell.font = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
title_cell.fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
title_cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 38

ws.merge_cells('A2:N2')
sub_cell = ws['A2']
sub_cell.value = "Iltimos, har bir xodim bo'yicha sariq rangli ustunlarni to'ldiring. Ushbu ma'lumotlar HR tabel, oylik hisob-kitob va tizim uchun ishlatiladi."
sub_cell.font = Font(name='Calibri', size=10, italic=True, color='334155')
sub_cell.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
sub_cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 24

headers = [
    ("№", 6),
    ("F.I.SH. (To'liq Familiya, Ism, Sharif)", 38),
    ("Rasmiy Lavozimi", 20),
    ("Tizimdagi Roli", 18),
    ("Telefon Raqami *", 20),
    ("Tug'ilgan Sana (KK.OO.YYYY) *", 18),
    ("Pasport Seriya va Raqam", 18),
    ("JSHSHIR (PINFL 14 xona)", 20),
    ("Oylik Oklad Miqdori (so'm)", 22),
    ("Ish Jadvali / Smena Turi", 24),
    ("Mutaxassisligi (Shifokorlar)", 25),
    ("Tibbiy Toifasi", 16),
    ("Tizimga Kirish Logini", 18),
    ("Qo'shimcha Izoh", 22)
]

header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

ws.row_dimensions[3].height = 36
for col_idx, (header_text, width) in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=header_text)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

# 16 Staff members from official list
staff_list = [
    (1, "ABDUKARIMOVA RA'NO XUDAYBERGANOVNA", "Медсестра", "Navbatchi hamshira", "nurse_rano", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "1-toifa"),
    (2, "ATAMETOVA TURSUNOY XUDAYBERGANOVNA", "Медсестра", "Navbatchi hamshira", "nurse_tursunoy", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "Oliy toifa"),
    (3, "DONOBOYEVA LAYLO DONOBOYEVNA", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (4, "EGAMQULOVA XILOLA XOSHIMQULOVNA", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (5, "G'ANIYEVA SAYGUL SHODMON QIZI", "Медсестра", "Navbatchi hamshira", "nurse_saygul", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "2-toifa"),
    (6, "IBRAGIMOVA NODIRA BATIROVNA", "Массажист", "Massajchi / Fizioterapiya", "spec_nodira", "Tibbiy massaj", "5 000 000", "Kunduzgi (09:00 - 17:00)", "Mutaxassis"),
    (7, "KARABAEVA MUXAYYO TOSHPO'LATOVNA", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (8, "KARAKULOVA XURSHIDA MIRZAKIM QIZI", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (9, "MAXKAMOVA UMIDA MUMINOVNA", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (10, "MUSAXANOVA FOTIMA DANABAYEVNA", "Санитарка", "Kichik tibbiy xodim", "sanitar", "Sanitariya & Gigiyena", "3 800 000", "24 soatlik navbatchilik", "Mutaxassis"),
    (11, "SAFAROVA SHAXNOZA ABDUMANNAPOVNA", "Медсестра", "Navbatchi hamshira", "nurse_shaxnoza", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "1-toifa"),
    (12, "SAYDAMINOVA ZIYEDA UMAR QIZI", "Медсестра", "Navbatchi hamshira", "nurse_ziyeda", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "2-toifa"),
    (13, "SHERMUXAMEDOVA FARIDA MIRAXATOVNA", "Лечащий врач", "Davolovchi shifokor", "dr_farida", "Narkologiya", "12 000 000", "Kunduzgi (08:00 - 18:00)", "Oliy toifa"),
    (14, "UMAROV XUSAN PAYZIBOYEVICH", "гл. Врач", "Bosh shifokor", "dr_xusan", "Narkolog-psixiatr", "15 000 000", "Kunduzgi (08:00 - 18:00)", "Oliy toifa"),
    (15, "VASINA YULIYA ALEKSANDROVNA", "Лечащий врач", "Davolovchi shifokor", "dr_yuliya", "Psixiatriya", "12 000 000", "Kunduzgi (08:00 - 18:00)", "1-toifa"),
    (16, "XOLTOYEVA SHAXLO SOBITJONOVNA", "Медсестра", "Navbatchi hamshira", "nurse_shaxlo", "Hamshiralik ishi", "6 000 000", "24 soatlik navbatchilik", "1-toifa")
]

zebra_even = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
zebra_odd = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

# Required fill to highlight columns for user
highlight_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # soft amber

font_main = Font(name='Calibri', size=11, color='0F172A')
font_bold = Font(name='Calibri', size=11, bold=True, color='0F172A')
font_mono = Font(name='Consolas', size=10.5, color='0369A1', bold=True)
font_sanitar = Font(name='Calibri', size=10, italic=True, color='475569')

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

for num, name, role_doc, role_sys, login, spec, def_sal, def_shift, toifa in staff_list:
    row_num = num + 3
    ws.row_dimensions[row_num].height = 26
    base_fill = zebra_odd if num % 2 == 1 else zebra_even

    login_font = font_sanitar if "umumiy" in login else font_mono

    row_data = [
        (num, align_center, font_bold, base_fill),
        (name, align_left, font_bold, base_fill),
        (role_doc, align_center, font_main, base_fill),
        (role_sys, align_center, font_main, base_fill),
        ("+998 ", align_center, font_main, highlight_fill), # phone
        ("", align_center, font_main, highlight_fill),      # birthdate
        ("", align_center, font_main, base_fill),           # passport
        ("", align_center, font_main, base_fill),           # pinfl
        (def_sal, align_right, font_main, base_fill),       # salary
        (def_shift, align_center, font_main, base_fill),    # shift
        (spec, align_left, font_main, base_fill),           # spec
        (toifa, align_center, font_main, base_fill),        # toifa
        (login, align_center, login_font, base_fill),       # login
        ("", align_left, font_main, base_fill)              # notes
    ]

    for col_idx, (val, al, fn, fl) in enumerate(row_data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=val)
        cell.fill = fl
        cell.alignment = al
        cell.font = fn
        cell.border = thin_border

# Sheet 2: Yo'riqnoma & Smenalar Tushuntirishi
ws2 = wb.create_sheet(title="Yo'riqnoma & Eslatmalar")
ws2.views.sheetView[0].showGridLines = True

instructions = [
    ("Ustun Nomi", "Tavsif va Maslahat", "Misol"),
    ("Telefon Raqami", "Xodimning shaxsiy yoki xizmat raqami", "+998 90 123 45 67"),
    ("Tug'ilgan Sana", "Kun.Oy.Yil formatida", "15.08.1985"),
    ("Pasport Seriya va Raqam", "Shaxsini tasdiqlovchi hujjat", "AA 1234567"),
    ("JSHSHIR (PINFL)", "14 xonali yagona identifikatsiya raqami", "31508851230045"),
    ("Asosiy Oylik Oklad", "Soliqlarsiz oylik asosiy stavka (so'mda)", "6 000 000"),
    ("Ish Jadvali / Smena", "Xodimning navbatchilik rejimi", "24 soatlik navbatchilik / Kunduzgi"),
    ("Tizimga Kirish (Login)", "Shifokor va hamshiralar uchun alohida, sanitarkalar uchun 1 ta umumiy login", "dr_xusan / nurse_rano / sanitar"),
]

ws2.merge_cells('A1:C1')
t2 = ws2['A1']
t2.value = "JADVALNI TO'LDIRISH BO'YICHA ESLATMA"
t2.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
t2.fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
t2.alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 32

ws2.column_dimensions['A'].width = 25
ws2.column_dimensions['B'].width = 45
ws2.column_dimensions['C'].width = 30

for r_idx, row in enumerate(instructions, 2):
    ws2.row_dimensions[r_idx].height = 24
    is_hdr = (r_idx == 2)
    for c_idx, val in enumerate(row, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        cell.border = thin_border
        if is_hdr:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
        else:
            cell.font = font_main
            cell.alignment = align_left if c_idx != 1 else align_center

# Save file in project directory
out_dir = r"d:\02.Areas\fayzcontrol\fayzcontrol.uz"
out_file = os.path.join(out_dir, "FMH_Xodimlar_Royxati_2026.xlsx")
wb.save(out_file)
print("Saved to:", out_file)

# Copy to user's Desktop
desktop_dir = r"C:\Users\Ituser\Desktop"
if os.path.exists(desktop_dir):
    desktop_file = os.path.join(desktop_dir, "FMH_Xodimlar_Royxati_2026.xlsx")
    wb.save(desktop_file)
    print("Saved to Desktop:", desktop_file)
